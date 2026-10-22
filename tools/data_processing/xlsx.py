"""Restricted Hermes XLSX adaptation for literal, inspectable local operations.

Adapted from NousResearch/hermes-agent, XLSX 1.1.0, xlsx_read.py and
xlsx_create.py (jsonable/sheet_rows, paired formula/cache views, new-book rows).
Copyright (c) 2026 Nous Research. MIT license retained in
docs/data_skills/licenses/HERMES_MIT.txt; provenance in SOURCE_MANIFEST.md.
No upstream command line entrypoints or in-place edit/recalc code are installed.
"""

import csv
from datetime import date, datetime, time
import io
from pathlib import Path
import posixpath
import re
import shutil
import zipfile
from xml.etree.ElementTree import SubElement, tostring

from defusedxml import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from .safe_output import write_new_bytes
from .tabular import read_csv


_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def jsonable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def sheet_rows(ws, **bounds):
    return [[jsonable(c) for c in row] for row in ws.iter_rows(values_only=True, **bounds)]


def _package(path, max_uncompressed_bytes):
    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx is supported; no macro-enabled or legacy workbooks")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (len(entries) > 10000 or sum(i.file_size for i in entries) > max_uncompressed_bytes
                    or len({i.filename for i in entries}) != len(entries)):
                raise ValueError("XLSX archive budget exceeded or duplicate members")
            xml = {}
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    raw = archive.read(entry.filename)
                    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                        raise ValueError("XML document types/entities are forbidden")
                    xml[entry.filename] = ET.fromstring(raw)
            workbook = xml["xl/workbook.xml"]
            relationships = {r.attrib["Id"]: r.attrib["Target"]
                             for r in xml["xl/_rels/workbook.xml.rels"]
                             if r.attrib.get("TargetMode") != "External"}
            sheets, extents, total_cells = {}, {}, 0
            for sheet in workbook.find(_MAIN + "sheets"):
                target = relationships[sheet.attrib[_REL + "id"]]
                member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
                if not member.startswith("xl/") or member not in xml:
                    raise ValueError("Worksheet target is missing or outside the workbook")
                cells, rows, columns = {}, 0, 0
                for cell in xml[member].iter(_MAIN + "c"):
                    coordinate = cell.attrib["r"]
                    if coordinate in cells:
                        raise ValueError("Duplicate worksheet cell coordinate")
                    r, c = coordinate_to_tuple(coordinate)
                    rows, columns = max(rows, r), max(columns, c)
                    if total_cells + rows * columns > 100000:
                        raise ValueError("XLSX actual cell extent budget exceeded")
                    cells[coordinate] = cell
                name = sheet.attrib["name"]
                sheets[name], extents[name] = cells, (rows, columns)
                total_cells += rows * columns
            return sheets, extents, any("externalLinks/" in i.filename or i.filename == "xl/connections.xml" for i in entries)
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ValueError("Invalid XLSX package") from exc


def inspect_xlsx(path, max_uncompressed_bytes=16 * 1024 * 1024) -> dict:
    if type(max_uncompressed_bytes) is not int or max_uncompressed_bytes <= 0:
        raise ValueError("XLSX archive budget must be a positive integer")
    lexical, extents, external = _package(path, max_uncompressed_bytes)
    wb_f = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    try:
        wb_v = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        try:
            sheets, count = [], 0
            for ws_f in wb_f.worksheets:
                ws_v = wb_v[ws_f.title]
                actual_rows, actual_columns = extents[ws_f.title]
                if (ws_f.max_row is not None and ws_f.max_row < actual_rows
                        or ws_f.max_column is not None and ws_f.max_column < actual_columns):
                    raise ValueError("Worksheet dimensions hide actual source cells")
                count += (ws_f.max_row or 1) * (ws_f.max_column or 1)
                if count > 100000:
                    raise ValueError("XLSX cell budget exceeded")
                bounds = {"max_row": actual_rows, "max_col": actual_columns}
                values = sheet_rows(ws_f, **bounds) if actual_rows else []
                cached_values = sheet_rows(ws_v, **bounds) if actual_rows else []
                cells = []
                for row_index, row in enumerate(ws_f.iter_rows(**bounds) if actual_rows else []):
                    output_row = []
                    for column_index, cell in enumerate(row):
                        coordinate = getattr(cell, "coordinate", None)
                        element = lexical[ws_f.title].get(coordinate)
                        node = None if element is None else element.find(_MAIN + "v")
                        xml_value = None if node is None else node.text
                        value = values[row_index][column_index]
                        formula = value if cell.data_type == "f" else None
                        cached = None
                        if formula is not None:
                            cached = xml_value if element.attrib.get("t", "n") == "n" else cached_values[row_index][column_index]
                        elif cell.data_type == "n" and xml_value is not None:
                            value = xml_value  # Numeric source of truth is XML text, never a float.
                        output_row.append({"coordinate": coordinate, "value": value,
                                           "data_type": cell.data_type, "number_format": cell.number_format,
                                           "formula": formula, "cached_value": cached,
                                           "cache_status": ("MISSING" if cached is None else "PRESENT") if formula is not None else "NOT_APPLICABLE",
                                           "xml_value": xml_value})
                    cells.append(output_row)
                # The upstream row helper supplies date serialization; override every
                # numeric/formula value from the exact cell record before returning it.
                values = [[cell["value"] for cell in row] for row in cells]
                sheets.append({"name": ws_f.title, "state": ws_f.sheet_state,
                               "max_row": actual_rows, "max_column": actual_columns,
                               "rows": values, "cells": cells})
            return {"sheets": sheets, "recalculation": "NOT_CHECKED",
                    "external_connections_present": external,
                    "external_connections_executed": False, "macros_executed": False}
        finally:
            wb_v.close()
    finally:
        wb_f.close()


def read_xlsx(path, sheet=None) -> dict:
    report = inspect_xlsx(path)
    if sheet is None and len(report["sheets"]) != 1:
        raise ValueError("Specify a sheet explicitly when the workbook has multiple sheets")
    selected = next((s for s in report["sheets"] if sheet is None or s["name"] == sheet), None)
    if selected is None:
        raise ValueError("Requested sheet does not exist")
    rows = selected["rows"]
    return {"format": "xlsx", "sheet": selected["name"], "sheet_state": selected["state"],
            "headers": ["" if v is None else str(v) for v in rows[0]] if rows else [],
            "rows": rows[1:], "source_row_numbers": list(range(2, len(rows) + 1)),
            "recalculation": report["recalculation"], "cells": selected["cells"]}


def write_xlsx(sheets: dict, output) -> dict:
    if not isinstance(sheets, dict) or not sheets:
        raise ValueError("A new workbook requires at least one named sheet")
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        if (not isinstance(name, str) or not name or len(name) > 31
                or re.search(r"[\\/*?:\[\]]", name) or not isinstance(rows, list)):
            raise ValueError("Invalid sheet name or rows")
        if name.casefold() in {n.casefold() for n in wb.sheetnames}:
            raise ValueError("Sheet names must be unique ignoring case")
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            if not isinstance(row, list):
                raise ValueError("Rows must be arrays")
            for c, value in enumerate(row, 1):
                if value is not None and not isinstance(value, str):
                    raise ValueError("New-book values must be literal strings or null; no implicit numeric conversion")
                if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
                    raise ValueError("Cell cannot be stored without loss")
                cell = ws.cell(r, c, value)
                if value is not None:
                    cell.data_type = "s"  # CSV text must never become a formula.
                    cell.number_format = "@"
    buffer = io.BytesIO()
    wb.save(buffer)
    # openpyxl omits <is> for literal '', which otherwise reopens as None.
    # Repair only newly generated inline strings, preserving the string/null distinction.
    fixed = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                xml = ET.fromstring(raw)
                for cell in xml.iter(_MAIN + "c"):
                    if cell.attrib.get("t") == "inlineStr" and cell.find(_MAIN + "is") is None:
                        inline = SubElement(cell, _MAIN + "is")
                        SubElement(inline, _MAIN + "t").text = ""
                raw = tostring(xml)
            target.writestr(item, raw)
    path = write_new_bytes(output, fixed.getvalue())
    return {"path": str(path), "sheets": wb.sheetnames, "recalculation": "NOT_CHECKED"}


def csv_to_xlsx(source, output, encoding=None, delimiter=None) -> dict:
    table = read_csv(source, encoding=encoding, delimiter=delimiter)
    return write_xlsx({"Data": [table["headers"], *table["rows"]]}, output)


def xlsx_to_csv(source, output, sheet=None, delimiter=",") -> dict:
    table = read_xlsx(source, sheet=sheet)
    buffer, escaped = io.StringIO(newline=""), []
    writer = csv.writer(buffer, delimiter=delimiter)
    for r, row in enumerate([table["headers"], *table["rows"]], 1):
        converted = []
        for c, value in enumerate(row, 1):
            text = "" if value is None else str(value)
            if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
                text = "'" + text
                escaped.append({"row": r, "column": c})
            converted.append(text)
        writer.writerow(converted)
    path = write_new_bytes(output, buffer.getvalue().encode("utf-8"))
    return {"path": str(path), "escaped_cells": escaped,
            "transformation": "Apostrophe prefix protects spreadsheet CSV imports; original XLSX unchanged"}


def require_recalculation(path):
    """Refuse a claim of recalculation. Engine presence is not calculation proof."""
    inspect_xlsx(path)
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        raise RuntimeError("RECALCULATION_UNAVAILABLE: no engine; formula text is retained")
    raise RuntimeError("RECALCULATION_NOT_PERFORMED: use a separately reviewed engine procedure on copies; cached values may be stale")
