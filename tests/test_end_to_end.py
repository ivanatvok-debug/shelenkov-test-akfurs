"""Invented four-file inputs; real Q1-Q4 code and local inference subprocess."""

import csv
from decimal import Decimal
from functools import partial
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
try:
    from akfurs import solver
except ImportError:
    solver = None

NAMES = ['01_crm_orders_export.csv', '02_bank_statement_rub.csv',
         '03_catalog_prices_reference.xlsx', '04_client_chats.zip']


class EndToEndTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(solver, 'Full CLI implementation is missing')
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.data = self.base / 'synthetic inputs'
        self.data.mkdir()
        self.rows = [
            ['A901', '1', '2025-02-01', 'Искра', 'SYNTH-X', '1', '17', 'RUB', '', 'Оплачен', '2026-09-01 12:00'],
            ['A902', '1', '2025-02-01', 'Искра', 'SYNTH-X', '1', '20', 'RUB', '', 'Оплачен', '2026-09-01 12:00']]
        self.payments = [['01.01.2025', '09:00', '0,01', 'RUR', 'Поступление', 'Invented payer', 'заказ A901'],
                         ['01.01.2025', '10:00', '9007199254740993,01', 'RUR', 'Поступление', 'Invented payer', 'без номера']]
        self.write_crm()
        self.write_bank()
        self.write_catalog()
        self.write_chat()
        self.fake = self.base / 'invented-model.py'
        self.fake.write_text('''import json,sys
from pathlib import Path
request=json.loads(sys.stdin.read().split('\\nINPUT_DATA\\n',1)[1])
events=[]
for msg in request['chat']['messages']:
 for kind,quote,qualification in [('manager_price','Цена изделия 20 рублей.','fixed'),('manager_invitation','Приезжайте на примерку.','none')]:
  if quote in msg['text']:
   events.append(dict(position=msg['position'],role=msg['role'],type=kind,quote=quote,basis='Invented local fixture',qualification=qualification))
answer=dict(semantic_status='COMPLETE',events=events,absence=dict(price='PRESENT',invitation='PRESENT'),uncertainties=[])
Path(sys.argv[sys.argv.index('--output-last-message')+1]).write_text(json.dumps(answer))
for record in [dict(type='thread.started'),dict(type='turn.started'),dict(type='turn.completed',usage=dict(input_tokens=1,output_tokens=1))]:
 print(json.dumps(record))
''', encoding='utf-8')
        self.runner = partial(solver.run_zip, command=[sys.executable, str(self.fake)], live=False)
        self.number = 0

    def write_csv(self, path, headers, rows, encoding):
        with path.open('w', encoding=encoding, newline='') as stream:
            writer = csv.writer(stream, delimiter=';')
            writer.writerow(headers)
            writer.writerows(rows)

    def write_crm(self):
        self.write_csv(self.data / NAMES[0],
                       ['Номер заказа', '№ строки', 'Дата создания', 'Менеджер', 'Артикул',
                        'Кол-во', 'Цена за ед.', 'Валюта', 'Скидка', 'Статус', 'Выгружено'], self.rows, 'utf-8-sig')

    def write_bank(self):
        self.write_csv(self.data / NAMES[1],
                       ['Дата операции', 'Время', 'Сумма', 'Валюта', 'Тип', 'Контрагент', 'Назначение платежа'],
                       self.payments, 'cp1251')

    def write_catalog(self, minimum=18):
        wb = Workbook()
        wb.active.title = 'Каталог'
        wb.active.append(['Артикул', 'Наименование'])
        wb.active.append(['SYNTH-X', 'Invented item'])
        wb.create_sheet('Переименования').append(['Старый артикул', 'Новый артикул', 'Дата перехода', 'Комментарий'])
        history = wb.create_sheet('Прайс_история')
        history.append(['Артикул', 'Действует с', 'Действует по', 'Рекомендованная цена, ₽', 'Минимально допустимая цена, ₽'])
        history.append(['SYNTH-X', '2025-01-01', None, 99, minimum])
        history.sheet_state = 'hidden'
        managers = wb.create_sheet('Менеджеры')
        managers.append(['Каноническое имя', 'Варианты написания в CRM', 'Дата выхода', 'Дата увольнения'])
        managers.append(['Искра Тестовая', 'Искра', '2025-01-01', None])
        wb.save(self.data / NAMES[2])

    def write_chat(self, reverse=False):
        text = ('Цена изделия 20 рублей. Приезжайте на примерку.' if reverse
                else 'Приезжайте на примерку. Цена изделия 20 рублей.')
        with zipfile.ZipFile(self.data / NAMES[3], 'w') as archive:
            archive.writestr('invented.txt', 'Чат SYNTH-A | invented\n[23.10.2026 10:00] Менеджер: ' + text + '\n')

    def calculate(self, **kwargs):
        self.number += 1
        evidence = self.base / f'run-{self.number}'
        with patch.object(solver, 'run_zip', self.runner):
            result = solver.calculate_all(self.data, evidence, run_command='python solve.py synthetic/', **kwargs)
        return result, evidence

    def test_complete_exact_json_and_evidence_without_input_changes(self):
        before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.data.iterdir()}
        answer, evidence = self.calculate()
        self.assertEqual([answer['q1_unique_orders'], answer['q1_unique_order_lines']], [2, 2])
        self.assertEqual(answer['q2a_paid_orders_without_payment'], [902])
        self.assertEqual(answer['q2b_unidentified_payments_count'], 1)
        self.assertEqual(answer['q2b_unidentified_payments_sum_rub'], Decimal('9007199254740993.01'))
        self.assertEqual(answer['q3_price_violation_lines_by_manager'], {'Искра Тестовая': 1})
        self.assertEqual(answer['q4_chats_meeting_before_price'], ['SYNTH-A'])
        self.assertEqual(len(answer), 10)
        self.assertTrue((evidence / 'Q1_EVIDENCE.json').is_file())
        self.assertTrue((evidence / 'Q2_RECONCILIATION.json').is_file())
        self.assertTrue((evidence / 'Q3_COVERAGE.json').is_file())
        self.assertTrue((evidence / 'q4' / 'REGISTRY.json').is_file())
        self.assertEqual(json.loads((evidence / 'RUN.json').read_text())['status'], 'PASS')
        self.assertEqual(before, {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.data.iterdir()})

    def test_repeat_recalculates_same_answer(self):
        first, _ = self.calculate()
        second, _ = self.calculate()
        self.assertEqual(first, second)

    def test_old_duplicate_does_not_change_counters(self):
        self.rows.append([*self.rows[0][:-1], '2026-08-01 12:00'])
        self.write_crm()
        answer, _ = self.calculate()
        self.assertEqual([answer['q1_unique_orders'], answer['q1_unique_order_lines']], [2, 2])

    def test_new_payment_changes_q2a(self):
        self.payments.append(['01.01.2025', '11:00', '0,01', 'RUR', 'Поступление', 'Invented payer', 'заказ A902'])
        self.write_bank()
        self.assertEqual(self.calculate()[0]['q2a_paid_orders_without_payment'], [])

    def test_changed_hidden_history_changes_q3(self):
        self.write_catalog(minimum=21)
        self.assertEqual(self.calculate()[0]['q3_price_violation_lines_total'], 2)

    def test_changed_text_order_changes_q4(self):
        self.write_chat(reverse=True)
        self.assertEqual(self.calculate()[0]['q4_chats_meeting_before_price'], [])

    def test_missing_each_input_prevents_model_calls(self):
        for name in NAMES:
            source = self.data / name
            absent = self.base / name
            source.rename(absent)
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.calculate()
            absent.rename(source)

    def test_held_crm_prevents_q4(self):
        self.rows.append([*self.rows[0][:6], '999', *self.rows[0][7:]])
        self.write_crm()
        with self.assertRaises(ValueError):
            self.calculate()
        self.assertFalse(list(self.base.glob('run-*/q4')))

    def test_q4_failure_does_not_create_answer(self):
        self.fake.write_text('import sys\nsys.exit(2)\n')
        with self.assertRaises(ValueError):
            self.calculate()
        self.assertFalse(list(self.base.glob('run-*/answers.json')))

    def test_final_evidence_failure_does_not_publish_answer(self):
        write = solver._write

        def fail_final(directory, name, value):
            if name == 'RUN.json':
                raise OSError('Invented final evidence fault')
            return write(directory, name, value)

        stdout, stderr = io.StringIO(), io.StringIO()
        from contextlib import redirect_stdout, redirect_stderr
        with patch.object(solver, 'run_zip', self.runner), patch.object(solver, '_write', fail_final), \
                redirect_stdout(stdout), redirect_stderr(stderr):
            code = solver.main([str(self.data), '--evidence', str(self.base / 'final-fault')])
        self.assertNotEqual(code, 0)
        self.assertEqual(stdout.getvalue(), '')
        self.assertFalse((self.base / 'final-fault' / 'answers.json').exists())

    def test_q4_failure_main_stdout_is_empty(self):
        self.fake.write_text('import sys\nsys.exit(2)\n')
        stdout, stderr = io.StringIO(), io.StringIO()
        from contextlib import redirect_stdout, redirect_stderr
        with patch.object(solver, 'run_zip', self.runner), redirect_stdout(stdout), redirect_stderr(stderr):
            code = solver.main([str(self.data), '--evidence', str(self.base / 'q4-fault')])
        self.assertNotEqual(code, 0)
        self.assertEqual(stdout.getvalue(), '')

    def test_existing_evidence_is_preserved(self):
        _, evidence = self.calculate()
        marker = evidence / 'sentinel'
        marker.write_text('preserve')
        with self.assertRaises((ValueError, FileExistsError)):
            solver.calculate_all(self.data, evidence, run_command='python solve.py synthetic/')
        self.assertEqual(marker.read_text(), 'preserve')

    def test_evidence_inside_inputs_is_rejected(self):
        with self.assertRaises(ValueError):
            solver.calculate_all(self.data, self.data / 'evidence', run_command='python solve.py synthetic/')
        self.assertFalse((self.data / 'evidence').exists())

    def test_cli_subprocess_only_full_json_in_stdout(self):
        bootstrap = (
            'import sys,runpy,functools;'
            f'sys.path.insert(0,{str(ROOT)!r});'
            f'sys.path.insert(0,{str(ROOT / "src")!r});'
            'from akfurs import solver;'
            f'solver.run_zip=functools.partial(solver.run_zip,command=[sys.executable,{str(self.fake)!r}],live=False);'
            f'sys.argv=[{str(ROOT / "solve.py")!r},{str(self.data)!r},"--evidence",{str(self.base / "cli")!r}];'
            f'runpy.run_path({str(ROOT / "solve.py")!r},run_name="__main__")')
        result = subprocess.run([sys.executable, '-c', bootstrap], cwd=self.base, capture_output=True,
                                text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'), timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        answer = json.loads(result.stdout, parse_float=Decimal)
        self.assertEqual(answer['q4_chats_meeting_before_price'], ['SYNTH-A'])
        self.assertEqual(answer['q2b_unidentified_payments_sum_rub'], Decimal('9007199254740993.01'))
        self.assertIn(str(self.data), answer['run_command'])

    def test_actual_cli_missing_input_empty_stdout(self):
        result = subprocess.run([sys.executable, str(ROOT / 'solve.py'), str(self.base / 'absent')],
                                capture_output=True, text=True, timeout=10)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(result.stdout, '')
        self.assertTrue(result.stderr.strip())


if __name__ == '__main__':
    unittest.main()
