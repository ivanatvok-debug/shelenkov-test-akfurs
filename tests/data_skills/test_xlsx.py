import hashlib
import io
import re
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook, load_workbook

from tools.data_processing.safe_output import OUTPUT_ROOT, write_new_bytes
from tools.data_processing.xlsx import inspect_xlsx, read_xlsx, write_xlsx, csv_to_xlsx, xlsx_to_csv, require_recalculation


class XlsxTests(unittest.TestCase):
    def test_empty_literal_string_and_null_remain_distinct(self):
        path = self.root / "empty-literal.xlsx"
        write_xlsx({"Data": [["id", "empty", "null"], ["001", "", None]]}, path)
        self.assertEqual(read_xlsx(path)["rows"], [["001", "", None]])

    def test_understated_dimension_cannot_hide_source_rows(self):
        source = self.workbook(); target = self.root / "understated.xlsx"
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w") as dst:
            for item in src.infolist():
                raw = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    raw = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:C1"', raw)
                dst.writestr(item, raw)
        with self.assertRaises(ValueError):
            read_xlsx(target, sheet="Visible")

    def test_missing_dimension_cannot_bypass_actual_cell_extent_budget(self):
        wb = Workbook(); wb.active["A1"] = "id"; wb.active["A100001"] = "001"
        buffer = io.BytesIO(); wb.save(buffer); source = self.root / "dimensionless.xlsx"
        with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(source, "w") as dst:
            for item in src.infolist():
                raw = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    raw = re.sub(rb'<dimension[^>]*/>', b'', raw)
                dst.writestr(item, raw)
        with self.assertRaises(ValueError):
            inspect_xlsx(source)

    def setUp(self):
        OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
        self.directory = tempfile.TemporaryDirectory(prefix="synthetic-xlsx-", dir=OUTPUT_ROOT)
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def workbook(self):
        wb = Workbook()
        wb.active.title = "Visible"
        wb.active.append(["id", "amount", "formula"])
        wb.active.append(["0001", "0.30", "=1+1"])
        hidden = wb.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden.append(["literal", "=-not-a-formula"])
        hidden["B1"].data_type = "s"
        source = self.root / "source.xlsx"
        wb.save(source)
        return source

    def test_T03_multi_sheet_and_hidden_sheet_are_inventoried(self):
        report = inspect_xlsx(self.workbook())
        self.assertEqual([(s["name"], s["state"]) for s in report["sheets"]], [("Visible", "visible"), ("Hidden", "hidden")])

    def test_T04_empty_formula_cache_does_not_erase_formula(self):
        report = inspect_xlsx(self.workbook())
        cell = report["sheets"][0]["cells"][1][2]
        self.assertEqual(cell["value"], "=1+1")
        self.assertEqual(cell["formula"], "=1+1")
        self.assertIsNone(cell["cached_value"])
        self.assertEqual(cell["cache_status"], "MISSING")
        self.assertEqual(report["recalculation"], "NOT_CHECKED")
        self.assertIsNone(report["sheets"][1]["cells"][0][1]["formula"])

    def test_T14_all_operations_leave_sources_unchanged(self):
        source = self.workbook(); before = hashlib.sha256(source.read_bytes()).hexdigest()
        inspect_xlsx(source); read_xlsx(source, sheet="Visible")
        xlsx_to_csv(source, self.root / "export.csv", sheet="Visible")
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_T16_missing_recalculation_tool_is_an_explicit_refusal(self):
        source = self.workbook()
        with patch("shutil.which", return_value=None):
            with self.assertRaises(RuntimeError):
                require_recalculation(source)
        self.assertEqual(inspect_xlsx(source)["recalculation"], "NOT_CHECKED")

    def test_csv_to_xlsx_writes_literal_cells_without_type_coercion(self):
        source = self.root / "synthetic.csv"
        source.write_text('id,amount,text\n0001,0.100000000000000001,=1+1\n0002,-0.20,@SUM(A1:A2)\n', encoding="utf-8")
        output = self.root / "output.xlsx"
        csv_to_xlsx(source, output, delimiter=",")
        wb = load_workbook(output, data_only=False)
        try:
            self.assertEqual(wb.active["A2"].value, "0001")
            self.assertEqual(wb.active["B2"].value, "0.100000000000000001")
            self.assertEqual(wb.active["C2"].data_type, "s")
            self.assertEqual(wb.active["C3"].data_type, "s")
        finally:
            wb.close()

    def test_safe_csv_export_explains_injection_escaping(self):
        source = self.workbook()
        output = self.root / "export.csv"
        report = xlsx_to_csv(source, output, sheet="Visible")
        self.assertIn("'=1+1", output.read_text())
        self.assertEqual(report["escaped_cells"], [{"row": 2, "column": 3}])

    def test_sheet_selection_is_required_when_multiple_exist(self):
        with self.assertRaises(ValueError):
            read_xlsx(self.workbook())

    def test_existing_output_and_outside_path_are_refused(self):
        output = self.root / "existing.json"; output.write_bytes(b"unchanged")
        with self.assertRaises(FileExistsError):
            write_new_bytes(output, b"replace")
        self.assertEqual(output.read_bytes(), b"unchanged")
        with self.assertRaises(ValueError):
            write_new_bytes(self.root.parents[2] / "escaped.txt", b"test")

    def test_symlink_output_path_is_refused(self):
        target = self.root / "target"; target.mkdir()
        link = self.root / "linked"; link.symlink_to(target, target_is_directory=True)
        with self.assertRaises(ValueError):
            write_new_bytes(link / "file.json", b"test")
        self.assertFalse((target / "file.json").exists())

    def test_long_cell_is_refused_instead_of_truncated(self):
        with self.assertRaises(ValueError):
            write_xlsx({"Data": [["x" * 32768]]}, self.root / "long.xlsx")
        self.assertFalse((self.root / "long.xlsx").exists())

    def test_numeric_xml_lexeme_is_not_replaced_by_a_float(self):
        wb = Workbook(); wb.active["A1"] = 1; buffer = io.BytesIO(); wb.save(buffer)
        source = self.root / "numeric.xlsx"
        with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(source, "w") as dst:
            for item in src.infolist():
                raw = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    raw = raw.replace(b"<v>1</v>", b"<v>12345678901234567890.12345</v>")
                dst.writestr(item, raw)
        cell = inspect_xlsx(source)["sheets"][0]["cells"][0][0]
        self.assertEqual(cell["value"], "12345678901234567890.12345")
        self.assertEqual(cell["number_format"], "General")

    def test_archive_budget_and_xml_entities_are_refused(self):
        source = self.workbook()
        with self.assertRaises(ValueError):
            inspect_xlsx(source, max_uncompressed_bytes=10)
        hostile = self.root / "hostile.xlsx"
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(hostile, "w") as dst:
            for item in src.infolist():
                raw = src.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    raw = b'<!DOCTYPE x [<!ENTITY e "danger">]>' + raw
                dst.writestr(item, raw)
        with self.assertRaises(ValueError):
            inspect_xlsx(hostile)


if __name__ == "__main__":
    unittest.main()
